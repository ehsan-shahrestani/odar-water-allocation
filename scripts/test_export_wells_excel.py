#!/usr/bin/env python3
"""Unit tests for export_wells_excel.py."""

import io
import sys
import tarfile
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import openpyxl

from export_wells_excel import (
    export_wells,
    format_iso_to_jalali,
    gregorian_to_jalali,
    sanitize_filename,
)


class TestExportWellsExcel(unittest.TestCase):
    def test_jalali_conversion(self):
        # 2024-03-20 -> 1403/01/01
        self.assertEqual(gregorian_to_jalali(2024, 3, 20), (1403, 1, 1))
        # 2026-10-30 -> 1405/08/08
        self.assertEqual(gregorian_to_jalali(2026, 10, 30), (1405, 8, 8))

        self.assertEqual(format_iso_to_jalali("2024-03-20"), "1403/01/01")
        self.assertEqual(format_iso_to_jalali("2024-03-20 14:30:00", include_time=True), "1403/01/01 14:30")
        self.assertEqual(format_iso_to_jalali(None), "-")

    def test_sanitize_filename(self):
        self.assertEqual(sanitize_filename("چاه شماره ۱"), "چاه شماره ۱")
        self.assertEqual(sanitize_filename("چاه/موتور:شرق*"), "چاه-موتور-شرق-")
        self.assertEqual(sanitize_filename(""), "چاه_بدون_نام")

    def create_mock_dump(self) -> str:
        """Create mock PostgreSQL COPY blocks with test data."""
        # IDs
        well1_id = "11111111-1111-1111-1111-111111111111"
        well2_id = "22222222-2222-2222-2222-222222222222"

        admin_id = "aaaaaaaa-aaaa-aaaa-aaaa-aaaaaaaaaaaa"
        farmer1_id = "bbbbbbbb-bbbb-bbbb-bbbb-bbbbbbbbbbbb"
        farmer2_id = "cccccccc-cccc-cccc-cccc-cccccccccccc"

        wy1_id = "33333333-3333-3333-3333-333333333333"
        wy2_id = "44444444-4444-4444-4444-444444444444"

        wf1_id = "55555555-5555-5555-5555-555555555555"
        wf2_id = "66666666-6666-6666-6666-666666666666"

        alloc1_id = "77777777-7777-7777-7777-777777777777"
        alloc2_id = "88888888-8888-8888-8888-888888888888"

        usage1_id = "99999999-9999-9999-9999-999999999999"
        usage2_id = "00000000-0000-0000-0000-000000000000"

        dump = f"""
COPY "public"."wells" ("id", "name", "description", "representative_id", "created_at") FROM stdin;
{well1_id}\tچاه ولیعصر\tتوضیحات چاه ۱\t{admin_id}\t2024-01-01 00:00:00+00
{well2_id}\tچاه مزرعه/شمالی\t\\N\t\\N\t2024-01-01 00:00:00+00
\\.

COPY "public"."profiles" ("id", "full_name", "phone", "role", "is_active", "created_at") FROM stdin;
{admin_id}\tمدیر سیستم\t09120000000\tadmin\tt\t2024-01-01 00:00:00+00
{farmer1_id}\tرضا اکبری\t09152404098\tfarmer\tt\t2024-01-02 00:00:00+00
{farmer2_id}\tحسین حسینی\t09905913852\tfarmer\tt\t2024-01-03 00:00:00+00
\\.

COPY "public"."water_years" ("id", "well_id", "description", "start_date", "end_date", "hours_per_share", "created_at") FROM stdin;
{wy1_id}\t{well1_id}\tسال آبی ۱۴۰۲-۱۴۰۳\t2023-09-23\t2024-09-22\t12.5\t2023-09-20 00:00:00+00
{wy2_id}\t{well2_id}\tسال آبی ۱۴۰۳-۱۴۰۴\t2024-09-23\t2025-09-22\t\\N\t2024-09-20 00:00:00+00
\\.

COPY "public"."well_farmers" ("id", "well_id", "farmer_id", "display_name", "created_at") FROM stdin;
{wf1_id}\t{well1_id}\t{farmer1_id}\tحاج رضا اکبری\t2024-01-05 00:00:00+00
{wf2_id}\t{well1_id}\t{farmer2_id}\t\\N\t2024-01-06 00:00:00+00
\\.

COPY "public"."water_allocations" ("id", "water_year_id", "well_farmer_id", "allocated_hours", "created_at") FROM stdin;
{alloc1_id}\t{wy1_id}\t{wf1_id}\t20.00\t2024-01-10 00:00:00+00
{alloc2_id}\t{wy1_id}\t{wf2_id}\t10.00\t2024-01-10 00:00:00+00
\\.

COPY "public"."water_usages" ("id", "allocation_id", "consumed_hours", "used_at", "description", "created_by", "created_at", "turned_on_pump", "received_from_well_farmer_id", "turned_off_pump", "delivered_to_well_farmer_id") FROM stdin;
{usage1_id}\t{alloc1_id}\t6.50\t2024-04-15 10:00:00+00\tآبیاری اول یونجه\t{admin_id}\t2024-04-15 16:30:00+00\tt\t\\N\tf\t{wf2_id}
{usage2_id}\t{alloc2_id}\t4.00\t2024-04-15 16:30:00+00\tتحویل از حاج رضا\t{admin_id}\t2024-04-15 20:30:00+00\tf\t{wf1_id}\tt\t\\N
\\.
"""
        return dump

    def test_export_wells_from_sql(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            tmppath = Path(tmpdir)
            sql_file = tmppath / "data.sql"
            sql_file.write_text(self.create_mock_dump(), encoding="utf-8")

            out_dir = tmppath / "output"
            exported = export_wells(sql_file, out_dir)

            self.assertEqual(len(exported), 2)

            # Check Well 1
            well1_file = out_dir / "چاه ولیعصر" / "چاه ولیعصر.xlsx"
            self.assertTrue(well1_file.is_file())

            # Check Well 2 (sanitized slash)
            well2_file = out_dir / "چاه مزرعه-شمالی" / "چاه مزرعه-شمالی.xlsx"
            self.assertTrue(well2_file.is_file())

            # Inspect Well 1 Workbook
            wb1 = openpyxl.load_workbook(well1_file)
            self.assertEqual(wb1.sheetnames, ["سال‌های آبی", "کشاورزان", "آبیاری‌ها"])

            # Verify RTL on all sheets
            for sname in wb1.sheetnames:
                ws = wb1[sname]
                self.assertTrue(
                    ws.views.sheetView[0].rightToLeft,
                    f"Sheet '{sname}' must be RTL",
                )
                self.assertIsNotNone(
                    ws.auto_filter.ref,
                    f"Sheet '{sname}' must have AutoFilter enabled",
                )
                self.assertEqual(
                    ws.freeze_panes,
                    "A2",
                    f"Sheet '{sname}' must freeze row 1",
                )

            # Check Sheet 2 (کشاورزان) calculations for Well 1
            ws_farmers = wb1["کشاورزان"]
            # Row 1 is header. Rows 2 & 3 are farmers.
            rows = list(ws_farmers.iter_rows(values_only=True))
            # Header check
            self.assertIn("نام کشاورز", rows[0])
            self.assertIn("سهمیه باقیمانده (ساعت)", rows[0])

            # Farmer 1: Haj Reza, allocated 20.0, consumed 6.5, remaining 13.5
            haj_reza = next(r for r in rows[1:] if "حاج رضا" in str(r[1]))
            self.assertEqual(haj_reza[4], 20.0)  # allocated
            self.assertEqual(haj_reza[5], 6.5)   # consumed
            self.assertEqual(haj_reza[6], 13.5)  # remaining

            # Check Sheet 3 (آبیاری‌ها) for Well 1
            ws_usages = wb1["آبیاری‌ها"]
            u_rows = list(ws_usages.iter_rows(values_only=True))
            self.assertIn("وضعیت شروع نوبت", u_rows[0])
            self.assertIn("وضعیت پایان نوبت", u_rows[0])

            # Usage 1 had turned_on_pump=t and delivered_to wf2 (حسین حسینی)
            row_u1 = next(r for r in u_rows[1:] if r[3] == "حاج رضا اکبری")
            self.assertEqual(row_u1[6], "روشن کردن چاه")
            self.assertEqual(row_u1[7], "تحویل به: حسین حسینی")

            # Usage 2 had received_from wf1 (حاج رضا) and turned_off_pump=t
            row_u2 = next(r for r in u_rows[1:] if r[3] == "حسین حسینی")
            self.assertEqual(row_u2[6], "تحویل از: حاج رضا اکبری")
            self.assertEqual(row_u2[7], "خاموش کردن چاه")

    def test_export_wells_from_tar_archive(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            tmppath = Path(tmpdir)
            archive_path = tmppath / "backup.tar.gz"

            dump_content = self.create_mock_dump().encode("utf-8")
            with tarfile.open(archive_path, "w:gz") as tar:
                ti = tarfile.TarInfo("data.sql")
                ti.size = len(dump_content)
                tar.addfile(ti, io.BytesIO(dump_content))

            out_dir = tmppath / "output"
            exported = export_wells(archive_path, out_dir)
            self.assertEqual(len(exported), 2)
            self.assertTrue((out_dir / "چاه ولیعصر" / "چاه ولیعصر.xlsx").is_file())


if __name__ == "__main__":
    unittest.main()
