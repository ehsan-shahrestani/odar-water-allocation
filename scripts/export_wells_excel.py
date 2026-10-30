#!/usr/bin/env python3
"""Export well records to individual RTL Excel workbooks per well.

Generates an Excel workbook (.xlsx) for each well containing:
1. سال‌های آبی (Water Years)
2. کشاورزان و سهمیه‌ها (Farmers & Quotas)
3. آبیاری‌ها (Irrigations / Water Usages with AutoFilter)

All sheets are set to Right-To-Left (RTL) mode.
Can read from:
- A Supabase backup .tar.gz archive
- A data.sql dump file
"""

from __future__ import annotations

import argparse
import hashlib
import os
import re
import sys
import tarfile
import uuid
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print(
        "Error: openpyxl is required. Install with: pip install openpyxl",
        file=sys.stderr,
    )
    sys.exit(1)


# ---------------------------------------------------------------------------
# Jalali Date Conversion (Pure Python, zero-dependency)
# ---------------------------------------------------------------------------

def gregorian_to_jalali(gy: int, gm: int, gd: int) -> tuple[int, int, int]:
    """Convert Gregorian date to Jalali (Solar Hijri) date."""
    g_d_m = [0, 31, 59, 90, 120, 151, 181, 212, 243, 273, 304, 334]
    if gy > 1600:
        jy = 979
        gy -= 1600
    else:
        jy = 0
        gy -= 621
    gy2 = (gm > 2) and (gy + 1) or gy
    days = (
        (365 * gy)
        + ((gy2 + 3) // 4)
        - ((gy2 + 99) // 100)
        + ((gy2 + 399) // 400)
        - 80
        + gd
        + g_d_m[gm - 1]
    )
    jy += 33 * (days // 12053)
    days %= 12053
    jy += 4 * (days // 1461)
    days %= 1461
    if days > 365:
        jy += (days - 1) // 365
        days = (days - 1) % 365
    if days < 186:
        jm = 1 + (days // 31)
        jd = 1 + (days % 31)
    else:
        jm = 7 + ((days - 186) // 30)
        jd = 1 + ((days - 186) % 30)
    return jy, jm, jd


def normalize_persian_text(val: str | None) -> str:
    """Normalize Arabic letters to Persian (ی, ک), fix spaces and punctuation."""
    if not val:
        return ""
    text = str(val).strip()
    text = text.replace("ي", "ی").replace("ى", "ی")
    text = text.replace("ك", "ک")
    text = text.replace("ة", "ه")
    text = text.replace("ۀ", "ه")
    text = re.sub(r"[ \t]+", " ", text)
    return text


def format_persian_phone(phone: str | None) -> str:
    """Format mobile number cleanly for Persian display (e.g. ۰۹۱۵-۲۴۰-۴۰۹۸)."""
    if not phone:
        return "-"
    digits = "".join(c for c in str(phone) if c.isdigit())
    if digits.startswith("98") and len(digits) == 12:
        digits = "0" + digits[2:]
    if len(digits) == 11 and digits.startswith("09"):
        formatted = f"{digits[:4]}-{digits[4:7]}-{digits[7:]}"
        return to_persian_digits(formatted)
    return to_persian_digits(phone)


def format_iso_to_jalali(iso_str: str | None, include_time: bool = False) -> str:
    """Format an ISO / Postgres date string into a Jalali date string with Persian digits."""
    if not iso_str:
        return "-"
    clean = iso_str.strip().replace("T", " ")
    date_part = clean.split()[0]
    time_part = ""
    if " " in clean and len(clean.split()) > 1:
        time_part = clean.split()[1][:5]  # HH:MM

    try:
        parts = [int(p) for p in date_part.split("-")]
        if len(parts) == 3:
            jy, jm, jd = gregorian_to_jalali(parts[0], parts[1], parts[2])
            res = f"{to_persian_digits(f'{jy:04d}')}/{to_persian_digits(f'{jm:02d}')}/{to_persian_digits(f'{jd:02d}')}"
            if include_time and time_part:
                return f"{res} ساعت {to_persian_digits(time_part)}"
            return res
    except (ValueError, IndexError):
        pass
    return to_persian_digits(iso_str)


def sanitize_filename(name: str) -> str:
    """Make string safe for folder/file names across OS environments."""
    cleaned = re.sub(r'[/\\:*?"<>|]', "-", name).strip()
    return cleaned if cleaned else "چاه_بدون_نام"


def to_persian_digits(val: int | str) -> str:
    """Convert ASCII digits to Persian digits."""
    return str(val).translate(str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹"))


def format_hours_to_human_fa(hours_val: float | int | str | None) -> str:
    """Format fractional hours as readable Persian string (e.g. '۳ ساعت و ۲۰ دقیقه')."""
    if hours_val is None or hours_val == "" or hours_val == "-":
        return "-"
    try:
        val = float(hours_val)
    except (ValueError, TypeError):
        return str(hours_val)

    if abs(val) < 0.0001:
        return "۰ ساعت"

    is_negative = val < 0
    abs_val = abs(val)
    total_minutes = int(round(abs_val * 60))
    h = total_minutes // 60
    m = total_minutes % 60

    h_fa = to_persian_digits(h)
    m_fa = to_persian_digits(m)

    prefix = "منفی " if is_negative else ""
    if h == 0 and m == 0:
        return "۰ ساعت"
    if m == 0:
        return f"{prefix}{h_fa} ساعت"
    if h == 0:
        return f"{prefix}{m_fa} دقیقه"
    return f"{prefix}{h_fa} ساعت و {m_fa} دقیقه"


# ---------------------------------------------------------------------------
# SQL COPY Block Parser
# ---------------------------------------------------------------------------

COPY_HEADER = re.compile(
    r'^COPY "(?P<schema>[a-z_]+)"\."(?P<table>[a-z_]+)" '
    r'\((?P<columns>"[a-z_]+"(?:, "[a-z_]+")*)\) FROM stdin;$'
)


def parse_copy_line(line: str, columns: tuple[str, ...]) -> dict[str, str | None]:
    """Parse a single tab-delimited COPY line into a dict."""
    parts = line.split("\t")
    row: dict[str, str | None] = {}
    for col, val in zip(columns, parts):
        if val == r"\N":
            row[col] = None
        else:
            val = (
                val.replace(r"\\", "\\")
                .replace(r"\n", "\n")
                .replace(r"\r", "\r")
                .replace(r"\t", "\t")
            )
            row[col] = val
    return row


def parse_sql_dump(sql_text: str) -> dict[str, list[dict[str, str | None]]]:
    """Extract public tables and their rows from PostgreSQL data.sql dump."""
    lines = sql_text.splitlines()
    tables: dict[str, list[dict[str, str | None]]] = {}
    index = 0

    while index < len(lines):
        header = COPY_HEADER.fullmatch(lines[index])
        if not header:
            index += 1
            continue

        schema = header.group("schema")
        table = header.group("table")
        columns = tuple(part.strip('"') for part in header.group("columns").split(", "))
        index += 1

        rows: list[dict[str, str | None]] = []
        while index < len(lines) and lines[index] != r"\.":
            rows.append(parse_copy_line(lines[index], columns))
            index += 1

        if schema == "public":
            tables[table] = rows

        index += 1

    return tables


def read_data_sql(source_path: Path) -> str:
    """Read data.sql either directly or extracted from a .tar.gz archive."""
    if source_path.is_dir():
        cand = source_path / "data.sql"
        if cand.is_file():
            return cand.read_text(encoding="utf-8")
        raise FileNotFoundError(f"data.sql not found in directory: {source_path}")

    if tarfile.is_tarfile(source_path):
        with tarfile.open(source_path, "r:*") as archive:
            member = archive.extractfile("data.sql")
            if not member:
                raise FileNotFoundError("data.sql is missing inside archive")
            return member.read().decode("utf-8")

    return source_path.read_text(encoding="utf-8")


# ---------------------------------------------------------------------------
# Excel Workbook Generator
# ---------------------------------------------------------------------------

# Color palette: Deep Teal (Persian Water) & soft cool slate zebra
HEADER_BG = "0F766E"       # Deep Teal
HEADER_TEXT = "FFFFFF"     # White
ZEBRA_BG = "F8FAFC"        # Very light cool slate
BORDER_COLOR = "CBD5E1"    # Clean slate border
FONT_FAMILY = "Tahoma"     # Gold standard font for Persian Excel support across all OS
STATUS_ACTIVE = "15803D"   # Forest Green
STATUS_EXHAUSTED = "DC2626"# Red
STATUS_NONE = "64748B"     # Slate Gray


def create_header_style() -> tuple[Font, PatternFill, Alignment, Border]:
    font = Font(name=FONT_FAMILY, size=10, bold=True, color=HEADER_TEXT)
    fill = PatternFill(start_color=HEADER_BG, end_color=HEADER_BG, fill_type="solid")
    alignment = Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=2)
    border = Border(
        left=Side(style="thin", color=BORDER_COLOR),
        right=Side(style="thin", color=BORDER_COLOR),
        top=Side(style="medium", color=HEADER_BG),
        bottom=Side(style="medium", color=HEADER_BG),
    )
    return font, fill, alignment, border


def create_data_border() -> Border:
    thin = Side(style="thin", color=BORDER_COLOR)
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def style_worksheet(ws: openpyxl.worksheet.worksheet.Worksheet) -> None:
    """Apply RTL layout, reading order, freeze panes, autofilter, and auto column widths."""
    # 1. Right-To-Left
    ws.views.sheetView[0].rightToLeft = True

    # 2. Freeze header row
    ws.freeze_panes = "A2"

    # 3. AutoFilter
    if ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = ws.dimensions

    # 4. Set Header Row Height & Data Rows
    ws.row_dimensions[1].height = 32
    for r in range(2, ws.max_row + 1):
        ws.row_dimensions[r].height = 24

    # 5. Auto adjust column widths with generous Persian character padding
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            char_len = sum(1.5 if ord(c) > 128 else 1.0 for c in val_str)
            if char_len > max_len:
                max_len = int(char_len)
        ws.column_dimensions[col_letter].width = max(max_len + 6, 16)


def build_well_workbook(
    well: dict[str, str | None],
    water_years: list[dict[str, str | None]],
    well_farmers: list[dict[str, str | None]],
    profiles_by_id: dict[str, dict[str, str | None]],
    allocations_by_well_farmer: dict[str, list[dict[str, str | None]]],
    usages_by_alloc_id: dict[str, list[dict[str, str | None]]],
    all_well_usages: list[dict[str, Any]],
) -> openpyxl.Workbook:
    """Build a styled, 3-sheet, RTL Excel workbook for a single well."""
    wb = openpyxl.Workbook()
    raw_name = (well.get("name") or "چاه").strip()
    wb.properties.title = f"گزارش و پشتیبان چاه {raw_name}"
    wb.properties.subject = "پشتیبان اکسل سامانه تخصیص آب اُدار"
    wb.properties.creator = "سامانه اُدار (Odar)"
    wb.properties.description = f"تولید شده به صورت خودکار در {format_iso_to_jalali(datetime.now().isoformat(), include_time=True)}"

    header_font, header_fill, header_align, header_border = create_header_style()
    data_border = create_data_border()
    zebra_fill = PatternFill(start_color=ZEBRA_BG, end_color=ZEBRA_BG, fill_type="solid")
    font_regular = Font(name=FONT_FAMILY, size=10)
    font_bold = Font(name=FONT_FAMILY, size=10, bold=True)

    # -----------------------------------------------------------------------
    # Sheet 1: سال‌های آبی
    # -----------------------------------------------------------------------
    ws_wy = wb.active
    ws_wy.title = "سال‌های آبی"

    wy_headers = [
        "ردیف",
        "عنوان سال آبی",
        "تاریخ شروع (شمسی)",
        "تاریخ پایان (شمسی)",
        "ساعت به ازای هر سهم (اعشاری)",
        "معادل به ساعت و دقیقه",
        "تاریخ شروع (میلادی)",
        "تاریخ پایان (میلادی)",
        "تاریخ ثبت (شمسی)",
    ]
    ws_wy.append(wy_headers)

    for i, h in enumerate(wy_headers, start=1):
        cell = ws_wy.cell(row=1, column=i)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = header_border

    # Sort water years by start_date descending
    sorted_wys = sorted(water_years, key=lambda x: x.get("start_date") or "", reverse=True)

    for idx, wy in enumerate(sorted_wys, start=1):
        row_num = idx + 1
        h_per_share = wy.get("hours_per_share")
        h_val = float(h_per_share) if h_per_share is not None else "-"

        row_data = [
            idx,
            normalize_persian_text(wy.get("description") or "-"),
            format_iso_to_jalali(wy.get("start_date")),
            format_iso_to_jalali(wy.get("end_date")),
            h_val,
            format_hours_to_human_fa(h_per_share),
            wy.get("start_date") or "-",
            wy.get("end_date") or "-",
            format_iso_to_jalali(wy.get("created_at"), include_time=True),
        ]
        ws_wy.append(row_data)

        # Styling data row
        is_even = idx % 2 == 0
        for col_idx in range(1, len(row_data) + 1):
            c = ws_wy.cell(row=row_num, column=col_idx)
            c.font = font_regular
            c.border = data_border
            if is_even:
                c.fill = zebra_fill

            if col_idx == 2:
                c.alignment = Alignment(horizontal="right", vertical="center", readingOrder=2)
            else:
                c.alignment = Alignment(horizontal="center", vertical="center", readingOrder=2)

            if col_idx == 5 and isinstance(h_val, float):
                c.number_format = "#,##0.00"
            elif col_idx == 6:
                c.font = font_bold

    style_worksheet(ws_wy)

    # -----------------------------------------------------------------------
    # Sheet 2: کشاورزان و سهمیه‌ها
    # -----------------------------------------------------------------------
    ws_farmers = wb.create_sheet(title="کشاورزان")

    farmer_headers = [
        "ردیف",
        "نام کشاورز",
        "شماره تماس",
        "سال آبی",
        "سهمیه تخصیص داده شده (ساعت)",
        "کل مصرف (ساعت)",
        "باقیمانده (ساعت)",
        "باقیمانده (ساعت و دقیقه)",
        "درصد مصرف",
        "وضعیت",
        "تاریخ عضویت در چاه (شمسی)",
    ]
    ws_farmers.append(farmer_headers)

    for i, h in enumerate(farmer_headers, start=1):
        cell = ws_farmers.cell(row=1, column=i)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = header_border

    # Prepare farmer quota rows
    # Map water years by id
    wy_map = {wy["id"]: wy for wy in water_years if wy.get("id")}

    farmer_rows = []
    for wf in well_farmers:
        wf_id = wf.get("id") or ""
        farmer_id = wf.get("farmer_id") or ""
        profile = profiles_by_id.get(farmer_id, {})

        raw_display = (wf.get("display_name") or "").strip() or (profile.get("full_name") or "کشاورز").strip()
        farmer_name = normalize_persian_text(raw_display)
        farmer_phone = format_persian_phone(profile.get("phone"))
        created_at_fa = format_iso_to_jalali(wf.get("created_at"))

        allocations = allocations_by_well_farmer.get(wf_id, [])

        if not allocations:
            # Farmer has no allocations yet; still list them so they are not lost
            farmer_rows.append({
                "name": farmer_name,
                "phone": farmer_phone,
                "wy_desc": "بدون تخصیص سال آبی",
                "allocated": 0.0,
                "consumed": 0.0,
                "remaining": 0.0,
                "percent": "-",
                "status": "بدون سهمیه",
                "created_at": created_at_fa,
            })
        else:
            for alloc in allocations:
                wy_id = alloc.get("water_year_id")
                raw_wy_desc = wy_map.get(wy_id, {}).get("description") or "سال آبی نامشخص"
                wy_desc = normalize_persian_text(raw_wy_desc)
                alloc_h = float(alloc.get("allocated_hours") or 0.0)

                # Sum consumed hours for this allocation
                alloc_usages = usages_by_alloc_id.get(alloc.get("id", ""), [])
                consumed_h = sum(float(u.get("consumed_hours") or 0.0) for u in alloc_usages)
                rem_h = round(alloc_h - consumed_h, 4)

                if alloc_h > 0:
                    pct = f"{to_persian_digits(round((consumed_h / alloc_h) * 100, 1))}٪"
                else:
                    pct = "۰٪"

                if alloc_h == 0:
                    status = "بدون سهمیه"
                elif rem_h <= 0:
                    status = "اتمام سهمیه"
                else:
                    status = "دارای اعتبار"

                farmer_rows.append({
                    "name": farmer_name,
                    "phone": farmer_phone,
                    "wy_desc": wy_desc,
                    "allocated": alloc_h,
                    "consumed": round(consumed_h, 2),
                    "remaining": round(rem_h, 2),
                    "percent": pct,
                    "status": status,
                    "created_at": created_at_fa,
                })

    # Sort farmers by name
    farmer_rows.sort(key=lambda x: (x["name"], x["wy_desc"]))

    for idx, f_row in enumerate(farmer_rows, start=1):
        row_num = idx + 1
        data = [
            idx,
            f_row["name"],
            f_row["phone"],
            f_row["wy_desc"],
            f_row["allocated"],
            f_row["consumed"],
            f_row["remaining"],
            format_hours_to_human_fa(f_row["remaining"]),
            f_row["percent"],
            f_row["status"],
            f_row["created_at"],
        ]
        ws_farmers.append(data)

        is_even = idx % 2 == 0
        for col_idx in range(1, len(data) + 1):
            c = ws_farmers.cell(row=row_num, column=col_idx)
            c.font = font_regular
            c.border = data_border
            if is_even:
                c.fill = zebra_fill

            if col_idx == 2:
                c.alignment = Alignment(horizontal="right", vertical="center", readingOrder=2)
            else:
                c.alignment = Alignment(horizontal="center", vertical="center", readingOrder=2)

            if col_idx in (5, 6, 7):
                c.number_format = "#,##0.00"
            elif col_idx == 8:
                # Highlight remaining quota if exhausted
                if f_row["remaining"] <= 0 and f_row["allocated"] > 0:
                    c.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=STATUS_EXHAUSTED)
                else:
                    c.font = font_bold
            elif col_idx == 10:
                # Status font color
                if f_row["status"] == "دارای اعتبار":
                    c.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=STATUS_ACTIVE)
                elif f_row["status"] == "اتمام سهمیه":
                    c.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=STATUS_EXHAUSTED)
                else:
                    c.font = Font(name=FONT_FAMILY, size=9.5, color=STATUS_NONE)

    style_worksheet(ws_farmers)

    # -----------------------------------------------------------------------
    # Sheet 3: آبیاری‌ها (نوبت‌های آب) با قابلیت فیلتر بر اساس کشاورز و سال آبی
    # -----------------------------------------------------------------------
    ws_usages = wb.create_sheet(title="آبیاری‌ها")

    usage_headers = [
        "ردیف",
        "تاریخ آبیاری (شمسی)",
        "ساعت",
        "نام کشاورز",
        "سال آبی",
        "مدت زمان مصرف (ساعت و دقیقه)",
        "مدت مصرف (ساعت اعشاری)",
        "وضعیت شروع نوبت",
        "وضعیت پایان نوبت",
        "توضیحات",
        "ثبت‌کننده",
        "تاریخ و زمان (میلادی)",
    ]
    ws_usages.append(usage_headers)

    for i, h in enumerate(usage_headers, start=1):
        cell = ws_usages.cell(row=1, column=i)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = header_border

    # Sort usages by used_at descending (latest first)
    sorted_usages = sorted(all_well_usages, key=lambda x: x.get("used_at") or "", reverse=True)

    for idx, u in enumerate(sorted_usages, start=1):
        row_num = idx + 1
        used_at = u.get("used_at") or ""
        date_fa = format_iso_to_jalali(used_at)
        time_part = ""
        if " " in used_at:
            time_part = used_at.split()[1][:5]
        elif "T" in used_at:
            time_part = used_at.split("T")[1][:5]

        duration_val = float(u.get("consumed_hours") or 0.0)

        # Start status
        if u.get("turned_on_pump") in ("t", "true", True):
            start_status = "روشن کردن چاه"
        elif u.get("received_from_name"):
            start_status = f"تحویل از: {normalize_persian_text(u['received_from_name'])}"
        else:
            start_status = "-"

        # End status
        if u.get("turned_off_pump") in ("t", "true", True):
            end_status = "خاموش کردن چاه"
        elif u.get("delivered_to_name"):
            end_status = f"تحویل به: {normalize_persian_text(u['delivered_to_name'])}"
        else:
            end_status = "-"

        u_row = [
            idx,
            date_fa,
            to_persian_digits(time_part) if time_part else "-",
            normalize_persian_text(u.get("farmer_name") or "-"),
            normalize_persian_text(u.get("water_year_desc") or "-"),
            format_hours_to_human_fa(duration_val),
            duration_val,
            start_status,
            end_status,
            normalize_persian_text(u.get("description") or "-"),
            normalize_persian_text(u.get("recorder_name") or "-"),
            used_at or "-",
        ]
        ws_usages.append(u_row)

        is_even = idx % 2 == 0
        for col_idx in range(1, len(u_row) + 1):
            c = ws_usages.cell(row=row_num, column=col_idx)
            c.font = font_regular
            c.border = data_border
            if is_even:
                c.fill = zebra_fill

            if col_idx in (4, 10):
                # Right aligned text (Farmer name, description)
                c.alignment = Alignment(horizontal="right", vertical="center", readingOrder=2)
            else:
                c.alignment = Alignment(horizontal="center", vertical="center", readingOrder=2)

            if col_idx == 6:
                # Human readable duration (bold)
                c.font = font_bold
            elif col_idx == 7:
                # Decimal hours (numeric for sum formulas)
                c.number_format = "#,##0.00"

    style_worksheet(ws_usages)

    return wb


# ---------------------------------------------------------------------------
# Main Extraction & Export Pipeline
# ---------------------------------------------------------------------------

def export_wells(source_path: Path, output_dir: Path) -> list[Path]:
    """Parse backup source and generate individual Excel workbooks per well."""
    print(f"Reading backup source: {source_path}...")
    sql_text = read_data_sql(source_path)
    tables = parse_sql_dump(sql_text)

    wells = tables.get("wells", [])
    profiles = tables.get("profiles", [])
    water_years = tables.get("water_years", [])
    well_farmers = tables.get("well_farmers", [])
    water_allocations = tables.get("water_allocations", [])
    water_usages = tables.get("water_usages", [])

    print(
        f"Loaded: {len(wells)} wells, {len(profiles)} profiles, "
        f"{len(water_years)} water years, {len(well_farmers)} farmer memberships, "
        f"{len(water_allocations)} allocations, {len(water_usages)} usages."
    )

    profiles_by_id = {p["id"]: p for p in profiles if p.get("id")}
    wy_by_id = {wy["id"]: wy for wy in water_years if wy.get("id")}
    alloc_by_id = {a["id"]: a for a in water_allocations if a.get("id")}

    # Map well_farmers by id
    wf_by_id = {wf["id"]: wf for wf in well_farmers if wf.get("id")}

    def get_farmer_display(wf_id: str | None) -> str:
        if not wf_id:
            return ""
        wf = wf_by_id.get(wf_id)
        if not wf:
            return ""
        name = (wf.get("display_name") or "").strip()
        if not name:
            p = profiles_by_id.get(wf.get("farmer_id") or "")
            name = (p.get("full_name") or "").strip() if p else ""
        return normalize_persian_text(name or "کشاورز")

    output_dir.mkdir(parents=True, exist_ok=True)
    exported_files: list[Path] = []

    for well in wells:
        well_id = well.get("id")
        if not well_id:
            continue

        raw_well_name = normalize_persian_text(well.get("name") or "چاه_بدون_نام")
        safe_name = sanitize_filename(raw_well_name)

        # Well-specific entities
        w_water_years = [wy for wy in water_years if wy.get("well_id") == well_id]
        w_well_farmers = [wf for wf in well_farmers if wf.get("well_id") == well_id]

        w_farmer_ids = {wf["id"] for wf in w_well_farmers if wf.get("id")}

        # Allocations for farmers in this well
        w_allocations = [
            a for a in water_allocations if a.get("well_farmer_id") in w_farmer_ids
        ]
        w_alloc_ids = {a["id"] for a in w_allocations if a.get("id")}

        # Group allocations by well_farmer_id
        allocations_by_wf: dict[str, list[dict[str, str | None]]] = {}
        for a in w_allocations:
            wfid = a.get("well_farmer_id") or ""
            allocations_by_wf.setdefault(wfid, []).append(a)

        # Group usages by allocation_id
        usages_by_alloc: dict[str, list[dict[str, str | None]]] = {}
        all_well_usages: list[dict[str, Any]] = []

        for u in water_usages:
            aid = u.get("allocation_id") or ""
            if aid in w_alloc_ids:
                usages_by_alloc.setdefault(aid, []).append(u)

                alloc = alloc_by_id.get(aid, {})
                wy_id = alloc.get("water_year_id")
                wy = wy_by_id.get(wy_id, {})
                wf_id = alloc.get("well_farmer_id")
                farmer_name = get_farmer_display(wf_id)

                recorder = profiles_by_id.get(u.get("created_by") or "")
                recorder_name = normalize_persian_text(recorder.get("full_name") if recorder else "-")

                all_well_usages.append({
                    "id": u.get("id"),
                    "used_at": u.get("used_at"),
                    "consumed_hours": u.get("consumed_hours"),
                    "description": u.get("description"),
                    "turned_on_pump": u.get("turned_on_pump"),
                    "turned_off_pump": u.get("turned_off_pump"),
                    "received_from_name": get_farmer_display(u.get("received_from_well_farmer_id")),
                    "delivered_to_name": get_farmer_display(u.get("delivered_to_well_farmer_id")),
                    "farmer_name": farmer_name,
                    "water_year_desc": wy.get("description") or "نامشخص",
                    "recorder_name": recorder_name,
                })

        # Build workbook
        wb = build_well_workbook(
            well=well,
            water_years=w_water_years,
            well_farmers=w_well_farmers,
            profiles_by_id=profiles_by_id,
            allocations_by_well_farmer=allocations_by_wf,
            usages_by_alloc_id=usages_by_alloc,
            all_well_usages=all_well_usages,
        )

        # Create target folder: <output_dir>/<well_name>/<well_name>.xlsx
        well_folder = output_dir / safe_name
        well_folder.mkdir(parents=True, exist_ok=True)
        excel_path = well_folder / f"{safe_name}.xlsx"

        wb.save(excel_path)
        exported_files.append(excel_path)
        print(f"Generated: {excel_path} ({len(w_well_farmers)} farmers, {len(all_well_usages)} irrigations)")

    return exported_files


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Export well records into separate RTL Excel files per well."
    )
    parser.add_argument(
        "source",
        type=Path,
        help="Path to backup .tar.gz archive or data.sql dump file",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path("wells_backup"),
        help="Target output directory (default: ./wells_backup)",
    )

    args = parser.parse_args()

    if not args.source.exists():
        print(f"Error: source not found: {args.source}", file=sys.stderr)
        sys.exit(1)

    try:
        exported = export_wells(args.source, args.output_dir)
        print(f"\nSuccessfully exported {len(exported)} well workbook(s) into {args.output_dir}.")
    except Exception as exc:
        print(f"Error during export: {exc}", file=sys.stderr)
        import traceback
        traceback.print_exc()
        sys.exit(1)


if __name__ == "__main__":
    main()
